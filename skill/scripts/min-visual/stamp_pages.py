# -*- coding: utf-8 -*-
"""Stamp brand logo top-right on raw page PNGs and build a full-page-image PPTX.

    stamp_pages.py <raw_pages_dir> <out_dir> [--brand 谷子] [--ppt-name PPT.pptx]

Writes <out_dir>/pages/*.png and <out_dir>/PPT.pptx. Logo path and size come from
config/品牌包/<brand>/风格.md rows 「Logo」 and 「Logo 位置」.
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

from PIL import Image

sys.path.insert(0, str(next(d for d in Path(__file__).resolve().parents if (d / "_lib" / "mini.py").is_file()) / "_lib"))
import mini  # noqa: E402

MIN_SIZE = (1920, 1080)


def fit_16x9(im: Image.Image) -> Image.Image:
    """Keep 2K native pixels. Only letterbox-upscale frames below 1920×1080."""
    im = im.convert("RGB")
    w, h = im.size
    if w >= MIN_SIZE[0] and h >= MIN_SIZE[1]:
        return im
    tw, th = MIN_SIZE
    ratio = min(tw / w, th / h)
    nw = max(1, int(w * ratio))
    nh = max(1, int(h * ratio))
    scaled = im.resize((nw, nh), Image.Resampling.LANCZOS)
    canvas = Image.new("RGB", (tw, th), scaled.getpixel((min(2, nw - 1), min(2, nh - 1))))
    canvas.paste(scaled, ((tw - nw) // 2, (th - nh) // 2))
    return canvas


def knock_black(src: Path) -> Image.Image:
    im = Image.open(src).convert("RGBA")
    px = im.load()
    w, h = im.size
    for y in range(h):
        for x in range(w):
            r, g, b, a = px[x, y]
            if r < 28 and g < 28 and b < 28:
                px[x, y] = (0, 0, 0, 0)
    return im


def stamp(page: Path, logo: Image.Image, dest: Path, width_ratio: float) -> None:
    canvas = fit_16x9(Image.open(page)).convert("RGBA")
    pw, ph = canvas.size
    target_w = max(int(pw * width_ratio), 180)
    scale = target_w / logo.width
    mark = logo.resize((target_w, max(1, int(logo.height * scale))), Image.Resampling.LANCZOS)
    x = pw - mark.width - int(pw * 0.028)
    y = int(ph * 0.035)
    canvas.alpha_composite(mark, (x, y))
    dest.parent.mkdir(parents=True, exist_ok=True)
    canvas.convert("RGB").save(dest, quality=95)


def build_pptx(image_dir: Path, output: Path) -> None:
    from lxml import etree
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches

    files = sorted(image_dir.glob("*.png"), key=lambda p: p.name.casefold())
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]
    for image_path in files:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image_path), Emu(0), Emu(0), width=deck.slide_width, height=deck.slide_height
        )
        sld = slide._element
        for child in list(sld):
            if child.tag == qn("p:transition"):
                sld.remove(child)
        trans = etree.Element(qn("p:transition"))
        trans.set("spd", "slow")
        etree.SubElement(trans, qn("p:fade"))
        sld.append(trans)
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("raw_dir", type=Path)
    parser.add_argument("out_dir", type=Path)
    parser.add_argument("--brand", default="谷子")
    parser.add_argument("--no-pptx", action="store_true")
    parser.add_argument("--ppt-name", default="PPT.pptx", help="PPT 文件名；目标被占用时可另存新版")
    args = parser.parse_args()

    style = mini.read_table(mini.brand_dir(args.brand) / "风格.md")
    logo_path = mini.resolve_path(style.get("Logo", ""))
    if not logo_path.is_file():
        raise SystemExit(f"风格.md 的 Logo 不存在: {logo_path}")
    m = re.search(r"(\d+(?:\.\d+)?)\s*%", style.get("Logo 位置", ""))
    ratio = float(m.group(1)) / 100 if m else 0.13

    raws = sorted(args.raw_dir.glob("*.png"), key=lambda p: p.name.casefold())
    if not raws:
        raise SystemExit(f"no png in {args.raw_dir}")
    logo = knock_black(logo_path)
    pages = args.out_dir / "pages"
    pages.mkdir(parents=True, exist_ok=True)
    for old in pages.glob("*.png"):
        old.unlink()
    for src in raws:
        stamp(src, logo, pages / src.name, ratio)
        print("stamped", src.name, flush=True)
    script_path = args.out_dir / "PPT大纲.md"
    if script_path.is_file():
        doc = mini.parse_script(script_path)
        pick = next((page for page in doc.pages if page.index == 2), None) or next(
            (page for page in doc.pages if page.index == 1), None
        )
        if pick:
            matches = sorted(pages.glob(f"{pick.num}-*.png"))
            if matches:
                shutil.copyfile(matches[0], args.out_dir / "代表页.png")
                print("representative", matches[0].name)
    if not args.no_pptx:
        ppt_path = args.out_dir / args.ppt_name
        build_pptx(pages, ppt_path)
        print("pptx", ppt_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
